"""Conservative OOXML converter: editable basics, rendered fallback for complex objects."""

import base64, io, os, shutil, subprocess, tempfile, html, zipfile
from pathlib import Path
from pptx import Presentation
from PIL import Image
import fitz
from .content import normalize


def convert(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if (
            len(archive.infolist()) > 10000
            or sum(x.file_size for x in archive.infolist()) > 150 * 1024 * 1024
        ):
            raise ValueError("PPTX 解压体积超过限制")
    with tempfile.TemporaryDirectory(prefix="quark-pptx-") as tmp:
        root = Path(tmp)
        source = root / "source.pptx"
        source.write_bytes(data)
        prs = Presentation(io.BytesIO(data))
        w = round(prs.slide_width / 9525)
        h = round(prs.slide_height / 9525)
        if not (320 <= w <= 2560 and 200 <= h <= 10000):
            raise ValueError("PPTX 页面尺寸超出渲染范围")
        if len(prs.slides) > 100:
            raise ValueError("每个模板最多 100 页")
        soffice = os.getenv("SOFFICE_PATH") or shutil.which("soffice")
        if not soffice:
            raise ValueError("缺少 LibreOffice，无法验证 PPTX 保真效果；请安装后重试")
        p = subprocess.run(
            [
                soffice,
                "-env:UserInstallation=file://" + str(root / "profile"),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                str(source),
            ],
            capture_output=True,
            timeout=150,
        )
        if p.returncode or not (root / "source.pdf").exists():
            raise ValueError("PPTX 参考渲染失败，未创建可发布模板")
        pdf = fitz.open(root / "source.pdf")
        # Render only the master/layout artwork into an immutable background.
        # Removing every slide-local object avoids duplicating editable text.
        background_prs = Presentation(io.BytesIO(data))
        for background_slide in background_prs.slides:
            for shape in list(background_slide.shapes):
                shape._element.getparent().remove(shape._element)
        background_file = root / "background.pptx"
        background_prs.save(background_file)
        background_result = subprocess.run(
            [
                soffice,
                "-env:UserInstallation=file://" + str(root / "background-profile"),
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                str(background_file),
            ],
            capture_output=True,
            timeout=150,
        )
        if background_result.returncode or not (root / "background.pdf").exists():
            raise ValueError("PPTX 母版背景渲染失败，未创建可发布模板")
        background_pdf = fitz.open(root / "background.pdf")
        documents = []
        diagnostics = []
        for ix, slide in enumerate(prs.slides):
            ref = pdf[ix].get_pixmap(
                matrix=fitz.Matrix(w / pdf[ix].rect.width, h / pdf[ix].rect.height),
                alpha=False,
            )
            reference = Image.open(io.BytesIO(ref.tobytes("png")))
            background_pix = background_pdf[ix].get_pixmap(
                matrix=fitz.Matrix(
                    w / background_pdf[ix].rect.width,
                    h / background_pdf[ix].rect.height,
                ),
                alpha=False,
            )
            background_uri = (
                "data:image/png;base64,"
                + base64.b64encode(background_pix.tobytes("png")).decode()
            )
            parts = [
                f'<img data-readonly="true" alt="PPTX 母版背景" style="position:absolute;inset:0;width:{w}px;height:{h}px;pointer-events:none" src="{background_uri}">'
            ]
            notes = [
                {
                    "source": "pptx",
                    "severity": "warning",
                    "message": f"第 {ix + 1} 页：母版与背景固定保留；请对照参考图核对继承字体、行距与文字溢出。",
                }
            ]
            refs = []

            def fallback(shape, reason):
                x, y, sw, sh = [
                    round(v / 9525)
                    for v in (shape.left, shape.top, shape.width, shape.height)
                ]
                crop = reference.crop(
                    (max(0, x), max(0, y), min(w, x + sw), min(h, y + sh))
                )
                buf = io.BytesIO()
                crop.save(buf, format="PNG")
                parts.append(
                    f'<img data-readonly="true" alt="{html.escape(reason)}" style="position:absolute;left:{x}px;top:{y}px;width:{sw}px;height:{sh}px" src="data:image/png;base64,{base64.b64encode(buf.getvalue()).decode()}">'
                )
                notes.append(
                    {
                        "source": "pptx",
                        "severity": "warning",
                        "message": f"第 {ix + 1} 页：{reason}，保留为图像，内部文字不可单独编辑",
                    }
                )

            bg = "#ffffff"
            try:
                bg = "#" + str(slide.background.fill.fore_color.rgb)
            except Exception:
                pass

            def is_complex(shape):
                if (
                    shape.rotation
                    or shape.shape_type in (6, 3, 19, 24)
                    or shape.has_table
                    or shape.has_chart
                ):
                    return True
                if shape.shape_type == 13:
                    return any(
                        getattr(shape, "crop_" + side, 0)
                        for side in ("left", "right", "top", "bottom")
                    )
                if shape.has_text_frame:
                    try:
                        return shape.shape_type == 1 and int(
                            shape.auto_shape_type
                        ) not in (1, 5)
                    except Exception:
                        return False
                return True

            shapes = list(slide.shapes)
            fixed = set(i for i, s in enumerate(shapes) if is_complex(s))

            def intersects(a, b):
                return (
                    a.left < b.left + b.width
                    and b.left < a.left + a.width
                    and a.top < b.top + b.height
                    and b.top < a.top + a.height
                )

            # Extend each complex group to all intersecting slide objects, transitively.
            groups = []
            while fixed:
                group = {fixed.pop()}
                changed = True
                while changed:
                    changed = False
                    for other in range(len(shapes)):
                        if (
                            other not in group
                            and shapes[other].left
                            < max(shapes[k].left + shapes[k].width for k in group)
                            and shapes[other].left + shapes[other].width
                            > min(shapes[k].left for k in group)
                            and shapes[other].top
                            < max(shapes[k].top + shapes[k].height for k in group)
                            and shapes[other].top + shapes[other].height
                            > min(shapes[k].top for k in group)
                        ):
                            group.add(other)
                            fixed.discard(other)
                            changed = True
                groups.append(group)
            grouped = {i for g in groups for i in g}
            group_end = {max(g): g for g in groups}
            for shape_index, shape in enumerate(shapes):
                x, y, sw, sh = [
                    round(v / 9525)
                    for v in (shape.left, shape.top, shape.width, shape.height)
                ]
                style = f"position:absolute;box-sizing:border-box;left:{x}px;top:{y}px;width:{sw}px;height:{sh}px;"
                refs.append(
                    {"shapeId": shape.shape_id, "name": shape.name, "slide": ix + 1}
                )
                if shape_index in grouped:
                    if shape_index in group_end:
                        group = group_end[shape_index]
                        left = min(shapes[k].left for k in group)
                        top = min(shapes[k].top for k in group)
                        right = max(shapes[k].left + shapes[k].width for k in group)
                        bottom = max(shapes[k].top + shapes[k].height for k in group)
                        from types import SimpleNamespace

                        bounds = SimpleNamespace(
                            left=left, top=top, width=right - left, height=bottom - top
                        )
                        fallback(
                            bounds,
                            "复杂图形及其相交区域（" + str(len(group)) + " 个对象）",
                        )
                    continue
                if (
                    shape.rotation
                    or shape.shape_type in (6, 3, 19, 24)
                    or shape.has_table
                    or shape.has_chart
                ):
                    fallback(shape, "组合、旋转或复杂图形")
                    continue
                if shape.shape_type == 13:
                    if any(
                        getattr(shape, "crop_" + s, 0)
                        for s in ("left", "right", "top", "bottom")
                    ):
                        fallback(shape, "裁剪图片")
                        continue
                    image = shape.image
                    parts.append(
                        f'<img data-source-shape="{shape.shape_id}" style="{style}" src="data:{image.content_type};base64,{base64.b64encode(image.blob).decode()}">'
                    )
                    continue
                if shape.has_text_frame:
                    # Non-rectangular decorative shapes retain their reference region rather than duplicating text.
                    try:
                        if shape.shape_type == 1 and int(shape.auto_shape_type) not in (
                            1,
                            5,
                        ):
                            fallback(shape, "非矩形形状")
                            continue
                    except Exception:
                        pass
                    try:
                        style += "background:#" + str(shape.fill.fore_color.rgb) + ";"
                    except Exception:
                        pass
                    tf = shape.text_frame
                    style += f"padding:{tf.margin_top / 9525}px {tf.margin_right / 9525}px {tf.margin_bottom / 9525}px {tf.margin_left / 9525}px;overflow:hidden;"
                    lines = []
                    for para in tf.paragraphs:
                        align = {1: "left", 2: "center", 3: "right", 4: "justify"}.get(
                            para.alignment, "left"
                        )
                        runs = []
                        for run in para.runs:
                            font = run.font
                            fs = f"font-size:{font.size.pt * 96 / 72 if font.size else 24}px;font-family:{html.escape(font.name or 'Arial')};"
                            if font.bold:
                                fs += "font-weight:700;"
                            if font.italic:
                                fs += "font-style:italic;"
                            try:
                                fs += "color:#" + str(font.color.rgb) + ";"
                            except Exception:
                                pass
                            runs.append(
                                f'<span style="{fs}">{html.escape(run.text)}</span>'
                            )
                        lines.append(
                            f'<p style="margin:0;text-align:{align};line-height:1.15">'
                            + ("".join(runs) or "<br>")
                            + "</p>"
                        )
                    parts.append(
                        f'<div data-source-shape="{shape.shape_id}" style="{style}">'
                        + "".join(lines)
                        + "</div>"
                    )
                else:
                    fallback(shape, "不支持的图形")
            # Source reference retained for human review; never overlay duplicate reference text.
            ref64 = base64.b64encode(ref.tobytes("png")).decode()
            doc = normalize(
                {
                    "html": f'<div style="position:relative;width:{w}px;height:{h}px;background:{bg}">'
                    + "".join(parts)
                    + "</div>",
                    "css": "",
                    "layoutMode": "fixed",
                    "sourceSize": {"width": w, "height": h},
                    "sourceReferences": refs
                    + [{"referenceImage": "data:image/png;base64," + ref64}],
                    "diagnostics": notes,
                }
            )
            documents.append(doc)
            diagnostics.extend(notes)
        return documents, diagnostics
