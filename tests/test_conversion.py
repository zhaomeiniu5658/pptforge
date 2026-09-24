import io, os, zipfile, base64
from pathlib import Path
import pytest
from PIL import Image
from quark.importers import html_import
from quark.content import tool
from quark.pptx_converter import convert


def test_zip_resources_are_archived_and_paths_guarded():
    image = io.BytesIO()
    Image.new("RGB", (12, 12), "blue").save(image, format="PNG")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr(
            "page.html",
            '<link rel="stylesheet" href="style.css"><h1>ZIP 页面</h1><img src="image.png">',
        )
        z.writestr("style.css", "h1{color:blue}")
        z.writestr("image.png", image.getvalue())
    d = html_import(buf.getvalue(), "package.zip")[0]
    assert "data:image/png;base64," in d["html"]
    assert "color:blue" in d["css"]
    assert not any(x["severity"] == "error" for x in d["diagnostics"])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../escape.html", "bad")
    with pytest.raises(ValueError, match="路径"):
        html_import(buf.getvalue(), "unsafe.zip")


def test_chromium_render_and_dimension_diff():
    doc = {
        "html": "<main><h1>临床方案验证</h1><button>查看</button></main>",
        "css": "main{height:550px;background:#eef4ff}h1{color:#2563eb}",
    }
    r = tool("/internal/render", {"document": doc, "width": 1000, "height": 600})
    assert not r["errors"]
    assert r["metrics"]["missingImages"] == 0
    Path("artifacts/render.png").write_bytes(base64.b64decode(r["png"]))
    same = tool("/internal/compare", {"target": r["png"], "actual": r["png"]})
    assert same["score"] == 1
    assert same["dimensionMismatch"] is False
    buf = io.BytesIO()
    Image.new("RGB", (900, 600), "white").save(buf, format="PNG")
    diff = tool(
        "/internal/compare",
        {"target": base64.b64encode(buf.getvalue()).decode(), "actual": r["png"]},
    )
    assert diff["dimensionMismatch"] is True
    assert diff["score"] < 1


def test_pptx_text_is_editable_and_complex_group_marked():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(1))
    run = text.text_frame.paragraphs[0].add_run()
    run.text = "Clinical protocol / editable text"
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(37, 99, 235)
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_shape(
        MSO_SHAPE.CHEVRON, Inches(0.5), Inches(2), Inches(4), Inches(1)
    )
    shape.text = "Complex group"
    buf = io.BytesIO()
    prs.save(buf)
    Path("artifacts/fixture.pptx").write_bytes(buf.getvalue())
    docs, diagnostics = convert(buf.getvalue())
    assert len(docs) == 1
    d = docs[0]
    assert "Clinical protocol / editable text" in d["html"]
    assert 'data-readonly="true"' in d["html"]
    assert d["layoutMode"] == "fixed"
    assert any("内部文字不可" in x["message"] for x in diagnostics)
    assert (
        "Complex group" not in d["html"]
    )  # embedded only in image, no duplicated text overlay
    rendered = tool("/internal/render", {"document": d, "width": 960, "height": 540})
    Path("artifacts/pptx-converted.png").write_bytes(base64.b64decode(rendered["png"]))
    Path("artifacts/pptx-reference.png").write_bytes(
        base64.b64decode(d["sourceReferences"][-1]["referenceImage"].split(",", 1)[1])
    )


def test_zip_export_deduplicates_binary_resources():
    from quark.exports import bundle

    picture = "data:image/png;base64," + base64.b64encode(b"image-data").decode()
    page = '<img src="' + picture + '"><img src="' + picture + '">'
    with zipfile.ZipFile(io.BytesIO(bundle(page, [page, page], {"draft": False}))) as z:
        assert len([n for n in z.namelist() if n.startswith("assets/")]) == 1
        assert b"../assets/" in z.read("pages/001.html")
        assert b"assets/" in z.read("index.html")


def test_nested_export_css_queries_and_metadata():
    image = io.BytesIO()
    Image.new('RGB', (12, 12), 'blue').save(image, format='PNG')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        z.writestr('__MACOSX/._page.html', 'invalid')
        z.writestr('export/pages/01/code.html', '<title>第一张</title><link rel="stylesheet" href="../../css/main.css?v=2"><img src="../../images/a%20b.png?v=1">')
        z.writestr('export/css/main.css', '@import "colors.css"; h1{background:url(../images/a%20b.png)}')
        z.writestr('export/css/colors.css', 'h1{color:blue}')
        z.writestr('export/images/a b.png', image.getvalue())
        z.writestr('export/pages/01/screen.png', image.getvalue())
        z.writestr('export/index.html', '<p>Duplicate navigation</p>')
    docs = html_import(buf.getvalue(), 'bundle.zip')
    assert len(docs) == 1
    assert 'data:image/png;base64,' in docs[0]['html']
    assert 'color:blue' in docs[0]['css']
    assert docs[0]['sourceReferences'][0]['title'] == '第一张'
    assert 'referenceImage' not in docs[0]['sourceReferences'][0]
    assert not docs[0]['diagnostics']


def test_tailwind_static_config_and_no_execution():
    html = '''<script src="https://cdn.tailwindcss.com"></script>
    <script>tailwind.config={theme:{extend:{colors:{brand:'#33384c'}}}}</script>
    <h1 class="bg-brand flex text-[32px]">可编辑标题</h1>'''
    doc = html_import(html.encode(), 'code.html')[0]
    assert '.bg-brand' in doc['css'] and '51 56 76' in doc['css']
    assert '32px' in doc['css'] and '可编辑标题' in doc['html']
    assert '<script' not in doc['html']
    with pytest.raises(ValueError, match='动态代码'):
        html_import(html.replace("'#33384c'", "(()=>{throw Error('executed')})()").encode(), 'code.html')


def test_unsupported_remote_resource_is_reported(monkeypatch):
    from quark import importers
    calls = []
    def fetch(url):
        calls.append(url)
        raise ValueError('unsupported host')
    monkeypatch.setattr(importers, 'remote_asset', fetch)
    docs = html_import(b'<img src="https://example.invalid/image.png">', 'code.html', archive_remote=True)
    assert calls == ['https://example.invalid/image.png']
    assert any(d['severity'] == 'error' for d in docs[0]['diagnostics'])
    with pytest.raises(ValueError):
        importers.safe_path('C:/escape.html')


def test_zip_duplicate_and_symlink_rejected():
    from quark.importers import archive_files
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w') as z:
        info = zipfile.ZipInfo('page.html')
        info.external_attr = 0o120777 << 16
        z.writestr(info, 'outside')
    with pytest.raises(ValueError, match='符号链接'):
        archive_files(buf.getvalue())
